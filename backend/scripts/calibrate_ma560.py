#!/usr/bin/env python3
"""MA560-Degression: Sample-Merkmale (anzahl_gepruefte_geraete) × 560-Faktura
(Isterlös) per EQ → €/Gerät über Mengenbänder + Degression-Diagnose (P2).

Lauf: ./venv/bin/python scripts/calibrate_ma560.py
"""
import json
import re
from pathlib import Path

import numpy as np
import openpyxl

SAMPLE = Path.home() / "Desktop/TUEV/_extracted/batch_MA560_sample_results.json"
FAKTURA = Path(__file__).parent.parent.parent / "sources/data/files/versand_export_24_06/560-WP_final.xlsx"


def main():
    d = json.load(open(SAMPLE))
    merk = {}
    for r in d:
        if not isinstance(r, dict) or r.get("error"):
            continue
        m = re.search(r"EQ(\d+)", r.get("file", ""))
        if not m:
            continue
        try:
            bm = int(r.get("anzahl_gepruefte_geraete"))
        except (TypeError, ValueError):
            bm = None
        merk[m.group(1)] = {"bm": bm, "typ": r.get("gebaeudetyp"), "branche": r.get("betreiber_branche")}

    wb = openpyxl.load_workbook(FAKTURA, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    price = {}
    for row in list(ws.iter_rows(values_only=True))[1:]:
        if row[3] is None:
            continue
        try:
            v = float(row[5]) if row[5] is not None else 0
        except (TypeError, ValueError):
            v = 0
        if v > 0:
            price[str(row[3])] = max(price.get(str(row[3]), 0), v)
    wb.close()

    pts = [(merk[eq]["bm"], price[eq]) for eq in (set(merk) & set(price))
           if merk[eq]["bm"] and merk[eq]["bm"] > 0 and 50 < price[eq] < 60000]
    print(f"Sample-EQ mit BM: {sum(1 for v in merk.values() if v['bm'])} · Faktura-EQ: {len(price)} · JOIN(BM+Preis): {len(pts)}")
    if len(pts) < 10:
        print("Zu wenige Join-Punkte — Sample evtl. noch nicht fertig.")
        return

    # €/Gerät über Mengenbänder
    bands = [(1, 50), (51, 150), (151, 400), (401, 800), (801, 99999)]
    print(f"\n{'Band BM':<14}{'n':>5}{'median €/BM':>13}{'median Preis':>14}")
    for lo, hi in bands:
        seg = [(bm, p) for bm, p in pts if lo <= bm <= hi]
        if not seg:
            continue
        rates = sorted(p / bm for bm, p in seg)
        prices = sorted(p for _, p in seg)
        print(f"{lo}-{hi if hi<99999 else '∞':<8}{len(seg):>5}{rates[len(rates)//2]:>13.2f}{prices[len(prices)//2]:>14.0f}")

    # einfaches Degression-Fit: log-log (Preis = a · BM^b)
    bm = np.array([x for x, _ in pts]); pr = np.array([y for _, y in pts])
    b, la = np.polyfit(np.log(bm), np.log(pr), 1)
    print(f"\nLog-Log-Fit:  Preis ≈ {np.exp(la):.0f} · BM^{b:.2f}   (b<1 = Degression)")
    # Vergleich aktuelles Modell 200 + 9,50·BM auf Testpunkten
    print(f"\nKontrolle Testfälle (aktuell 200+9,50·BM):")
    for eq, name in [("2229803", "Landgericht 850 BM"), ("2597657", "Auto Service 114 BM")]:
        if eq in price and eq in merk and merk[eq]["bm"]:
            b0 = merk[eq]["bm"]; real = price[eq]
            print(f"  {name}: real {real:.0f} ({real/b0:.2f}€/BM) · Modell {200+b0*9.5:.0f} ({(200+b0*9.5-real)/real*100:+.0f}%)")


if __name__ == "__main__":
    main()
